from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Initialize Presentation
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

# Theme Colors
NAVY = RGBColor(30, 41, 59)       # #1E293B
TEAL = RGBColor(14, 165, 233)     # #0EA5E9
SLATE = RGBColor(71, 85, 105)     # #475569
BG_LIGHT = RGBColor(248, 250, 252) # #F8FAFC
WHITE = RGBColor(255, 255, 255)
CARD_BG = RGBColor(255, 255, 255)
CARD_BORDER = RGBColor(226, 232, 240)
ACCENT_GREEN = RGBColor(16, 185, 129)
ACCENT_PURPLE = RGBColor(139, 92, 246)
TEXT_DARK = RGBColor(15, 23, 42)

def set_slide_background(slide, color=BG_LIGHT):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = color
    bg.line.color.rgb = color
    return bg

def add_header(slide, title_text, category_text="JOB PORTAL PROJECT"):
    header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(1.0))
    tf = header_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    
    p0 = tf.paragraphs[0]
    p0.text = category_text.upper()
    p0.font.name = 'Arial'
    p0.font.size = Pt(10)
    p0.font.bold = True
    p0.font.color.rgb = TEAL
    
    p1 = tf.add_paragraph()
    p1.text = title_text
    p1.font.name = 'Arial'
    p1.font.size = Pt(26)
    p1.font.bold = True
    p1.font.color.rgb = NAVY

def add_card(slide, left, top, width, height, bg_color=WHITE, border_color=CARD_BORDER):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    if border_color:
        card.line.color.rgb = border_color
        card.line.width = Pt(1)
    else:
        card.line.fill.background()
    return card

# -----------------------------------------------------------------------------
# SLIDE 1: TITLE SLIDE
# -----------------------------------------------------------------------------
slide_1 = prs.slides.add_slide(blank_layout)
set_slide_background(slide_1, NAVY)

# Left Color Accent Bar
accent_bar = slide_1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.5), Inches(0.15), Inches(4.5))
accent_bar.fill.solid()
accent_bar.fill.fore_color.rgb = TEAL
accent_bar.line.fill.background()

# Title Content Box
title_box = slide_1.shapes.add_textbox(Inches(1.2), Inches(1.5), Inches(11.0), Inches(4.5))
tf1 = title_box.text_frame
tf1.word_wrap = True

p = tf1.paragraphs[0]
p.text = "JOB PORTAL"
p.font.name = 'Arial'
p.font.size = Pt(44)
p.font.bold = True
p.font.color.rgb = WHITE

p2 = tf1.add_paragraph()
p2.text = "Online Job Searching and Recruitment System"
p2.font.name = 'Arial'
p2.font.size = Pt(22)
p2.font.color.rgb = TEAL
p2.space_after = Pt(40)

p3 = tf1.add_paragraph()
p3.text = "Presented By: Saksham | BCA – 3rd Semester"
p3.font.name = 'Arial'
p3.font.size = Pt(16)
p3.font.bold = True
p3.font.color.rgb = WHITE

p4 = tf1.add_paragraph()
p4.text = "Subject: Web Development"
p4.font.name = 'Arial'
p4.font.size = Pt(14)
p4.font.color.rgb = RGBColor(148, 163, 184)
p4.space_after = Pt(20)

p5 = tf1.add_paragraph()
p5.text = "Technologies Used: HTML  |  CSS  |  JavaScript  |  SQL"
p5.font.name = 'Arial'
p5.font.size = Pt(14)
p5.font.bold = True
p5.font.color.rgb = TEAL

# -----------------------------------------------------------------------------
# SLIDE 2: INTRODUCTION
# -----------------------------------------------------------------------------
slide_2 = prs.slides.add_slide(blank_layout)
set_slide_background(slide_2)
add_header(slide_2, "What is a Job Portal?")

# Left Box: Definition & Main Idea
left_card = add_card(slide_2, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8), NAVY, None)
tb_left = slide_2.shapes.add_textbox(Inches(1.1), Inches(2.1), Inches(5.0), Inches(4.2))
tf = tb_left.text_frame
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "OVERVIEW"
p.font.name = 'Arial'
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = TEAL
p.space_after = Pt(14)

p = tf.add_paragraph()
p.text = "A Job Portal is a web-based platform that connects job seekers with employers, streamlining the hiring process for both parties."
p.font.size = Pt(16)
p.font.color.rgb = WHITE
p.space_after = Pt(28)

p = tf.add_paragraph()
p.text = "MAIN IDEA"
p.font.name = 'Arial'
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = TEAL
p.space_after = Pt(10)

p = tf.add_paragraph()
p.text = "To make job searching and recruitment faster, easier, and highly organized through a centralized digital solution."
p.font.size = Pt(15)
p.font.color.rgb = RGBColor(226, 232, 240)

# Right Top Card: Job Seekers
add_card(slide_2, Inches(6.8), Inches(1.8), Inches(5.7), Inches(2.25))
tb = slide_2.shapes.add_textbox(Inches(7.1), Inches(2.0), Inches(5.1), Inches(1.85))
tf = tb.text_frame
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "For Job Seekers"
p.font.bold = True
p.font.size = Pt(18)
p.font.color.rgb = NAVY
p.space_after = Pt(10)

for feat in ["Create a personal account & profile", "Search and filter job vacancies", "View detailed job descriptions & apply"]:
    p = tf.add_paragraph()
    p.text = f"•  {feat}"
    p.font.size = Pt(13)
    p.font.color.rgb = SLATE

# Right Bottom Card: Employers
add_card(slide_2, Inches(6.8), Inches(4.35), Inches(5.7), Inches(2.25))
tb = slide_2.shapes.add_textbox(Inches(7.1), Inches(4.55), Inches(5.1), Inches(1.85))
tf = tb.text_frame
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "For Employers"
p.font.bold = True
p.font.size = Pt(18)
p.font.color.rgb = NAVY
p.space_after = Pt(10)

for feat in ["Post new job vacancies with requirements", "View incoming candidate applications", "Manage and shortlist potential candidates"]:
    p = tf.add_paragraph()
    p.text = f"•  {feat}"
    p.font.size = Pt(13)
    p.font.color.rgb = SLATE

# -----------------------------------------------------------------------------
# SLIDE 3: PROBLEM STATEMENT & OBJECTIVES
# -----------------------------------------------------------------------------
slide_3 = prs.slides.add_slide(blank_layout)
set_slide_background(slide_3)
add_header(slide_3, "Problem Statement & Project Objectives")

# Left Column: Problems
add_card(slide_3, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
tb = slide_3.shapes.add_textbox(Inches(1.1), Inches(2.1), Inches(5.0), Inches(4.2))
tf = tb.text_frame
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "Traditional Hiring Challenges"
p.font.bold = True
p.font.size = Pt(20)
p.font.color.rgb = RGBColor(225, 29, 72)
p.space_after = Pt(16)

probs = [
    "Time-Consuming: Manual job searching takes significant effort.",
    "Fragmented Sources: Job listings are scattered across multiple platforms.",
    "Application Tracking: Hard for job seekers to track submission status.",
    "Employer Overhead: Difficult for companies to manage paper/email resumes."
]
for pr in probs:
    p = tf.add_paragraph()
    p.text = f"•  {pr}"
    p.font.size = Pt(13)
    p.font.color.rgb = SLATE
    p.space_after = Pt(10)

# Right Column: Objectives
add_card(slide_3, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8))
tb = slide_3.shapes.add_textbox(Inches(7.1), Inches(2.1), Inches(5.1), Inches(4.2))
tf = tb.text_frame
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "Project Objectives"
p.font.bold = True
p.font.size = Pt(20)
p.font.color.rgb = TEAL
p.space_after = Pt(16)

objs = [
    "User-Friendly Platform: Easy online job searching and recruitment.",
    "Smart Search: Allow users to search and filter suitable jobs.",
    "Direct Application: Enable job seekers to apply online instantly.",
    "Employer Portal: Provide tools for posting and managing job vacancies.",
    "SQL Storage: Store user, job, and application data securely.",
    "Efficiency: Drastically reduce recruitment time and effort."
]
for obj in objs:
    p = tf.add_paragraph()
    p.text = f"✔  {obj}"
    p.font.size = Pt(13)
    p.font.color.rgb = SLATE
    p.space_after = Pt(8)

# -----------------------------------------------------------------------------
# SLIDE 4: TECHNOLOGIES USED
# -----------------------------------------------------------------------------
slide_4 = prs.slides.add_slide(blank_layout)
set_slide_background(slide_4)
add_header(slide_4, "Technology Stack")

techs = [
    ("HTML", "Structure", "Used to create the core structure, markup, and page content for the web portal.", TEAL),
    ("CSS", "Styling & Layout", "Handles visual styling, colors, modern layout designs, buttons, and responsive pages.", ACCENT_PURPLE),
    ("JavaScript", "Interactivity", "Provides dynamic features, form validation, client-side filtering, and interactive UI.", ACCENT_GREEN),
    ("SQL", "Database", "Manages relational database tables to store users, posted jobs, and application data.", NAVY)
]

col_w = Inches(2.7)
gap = Inches(0.3)
left_base = Inches(0.8)

for i, (tech, role, desc, accent) in enumerate(techs):
    x = left_base + i * (col_w + gap)
    card = add_card(slide_4, x, Inches(1.8), col_w, Inches(4.8))
    
    # Top color stripe
    stripe = slide_4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.8), col_w, Inches(0.2))
    stripe.fill.solid()
    stripe.fill.fore_color.rgb = accent
    stripe.line.fill.background()
    
    tb = slide_4.shapes.add_textbox(x + Inches(0.2), Inches(2.2), col_w - Inches(0.4), Inches(4.2))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = tech
    p.font.bold = True
    p.font.size = Pt(24)
    p.font.color.rgb = NAVY
    
    p_sub = tf.add_paragraph()
    p_sub.text = role.upper()
    p_sub.font.bold = True
    p_sub.font.size = Pt(10)
    p_sub.font.color.rgb = accent
    p_sub.space_after = Pt(18)
    
    p_desc = tf.add_paragraph()
    p_desc.text = desc
    p_desc.font.size = Pt(13)
    p_desc.font.color.rgb = SLATE

# -----------------------------------------------------------------------------
# SLIDE 5: MAIN MODULES
# -----------------------------------------------------------------------------
slide_5 = prs.slides.add_slide(blank_layout)
set_slide_background(slide_5)
add_header(slide_5, "System Architecture & Modules")

modules = [
    ("Job Seeker Module", TEAL, [
        "Registration & Secure Login",
        "Profile Creation & Management",
        "Job Search with Filtering",
        "View Full Job Details",
        "One-Click Job Application"
    ]),
    ("Employer Module", ACCENT_PURPLE, [
        "Employer Account Login",
        "Post New Job Vacancies",
        "Manage Active Job Listings",
        "View Received Applications",
        "Review & Candidate Shortlisting"
    ]),
    ("Admin Module", NAVY, [
        "Manage All User Accounts",
        "Manage Employer Profiles",
        "Oversee Job Listings",
        "Monitor System Activity",
        "Data Maintenance & Control"
    ])
]

col_w = Inches(3.7)
gap = Inches(0.3)
left_base = Inches(0.8)

for i, (mod_title, color, items) in enumerate(modules):
    x = left_base + i * (col_w + gap)
    card = add_card(slide_5, x, Inches(1.8), col_w, Inches(4.8))
    
    # Header area inside card
    header_bg = slide_5.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, Inches(1.8), col_w, Inches(0.8))
    header_bg.fill.solid()
    header_bg.fill.fore_color.rgb = color
    header_bg.line.fill.background()
    
    tb_h = slide_5.shapes.add_textbox(x, Inches(1.95), col_w, Inches(0.5))
    tf_h = tb_h.text_frame
    p_h = tf_h.paragraphs[0]
    p_h.text = mod_title
    p_h.alignment = PP_ALIGN.CENTER
    p_h.font.bold = True
    p_h.font.size = Pt(16)
    p_h.font.color.rgb = WHITE
    
    tb = slide_5.shapes.add_textbox(x + Inches(0.3), Inches(2.8), col_w - Inches(0.6), Inches(3.6))
    tf = tb.text_frame
    tf.word_wrap = True
    
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = f"•  {item}"
        p.font.size = Pt(13)
        p.font.color.rgb = SLATE
        p.space_after = Pt(12)

# -----------------------------------------------------------------------------
# SLIDE 6: HOW THE JOB PORTAL WORKS
# -----------------------------------------------------------------------------
slide_6 = prs.slides.add_slide(blank_layout)
set_slide_background(slide_6)
add_header(slide_6, "System Workflow & Process")

steps = [
    ("01", "Register / Login"),
    ("02", "Create Profile"),
    ("03", "Search Jobs"),
    ("04", "View Job Details"),
    ("05", "Apply for Job"),
    ("06", "Saved in SQL"),
    ("07", "Employer Review"),
    ("08", "Accept / Reject")
]

# 2 rows of 4 steps
box_w = Inches(2.6)
box_h = Inches(1.8)
gap_x = Inches(0.43)
gap_y = Inches(0.5)

for i, (num, label) in enumerate(steps):
    row = i // 4
    col = i % 4
    
    x = Inches(0.8) + col * (box_w + gap_x)
    y = Inches(1.9) + row * (box_h + gap_y)
    
    card = add_card(slide_6, x, y, box_w, box_h)
    
    tb = slide_6.shapes.add_textbox(x + Inches(0.15), y + Inches(0.15), box_w - Inches(0.3), box_h - Inches(0.3))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = num
    p.font.bold = True
    p.font.size = Pt(22)
    p.font.color.rgb = TEAL
    p.space_after = Pt(4)
    
    p2 = tf.add_paragraph()
    p2.text = label
    p2.font.bold = True
    p2.font.size = Pt(14)
    p2.font.color.rgb = NAVY

# -----------------------------------------------------------------------------
# SLIDE 7: DATABASE / SQL
# -----------------------------------------------------------------------------
slide_7 = prs.slides.add_slide(blank_layout)
set_slide_background(slide_7)
add_header(slide_7, "Database Structure (SQL)")

tables = [
    ("Users Table", TEAL, ["User_ID (PK)", "Name", "Email", "Password"]),
    ("Jobs Table", ACCENT_PURPLE, ["Job_ID (PK)", "Job_Title", "Company", "Location", "Salary"]),
    ("Applications Table", ACCENT_GREEN, ["Application_ID (PK)", "User_ID (FK)", "Job_ID (FK)", "Status"])
]

col_w = Inches(3.7)
gap = Inches(0.3)
left_base = Inches(0.8)

for i, (tname, color, fields) in enumerate(tables):
    x = left_base + i * (col_w + gap)
    card = add_card(slide_7, x, Inches(1.8), col_w, Inches(3.6))
    
    # Table Header Box
    hbox = slide_7.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, Inches(1.8), col_w, Inches(0.7))
    hbox.fill.solid()
    hbox.fill.fore_color.rgb = color
    hbox.line.fill.background()
    
    tb_h = slide_7.shapes.add_textbox(x, Inches(1.9), col_w, Inches(0.5))
    tf_h = tb_h.text_frame
    p_h = tf_h.paragraphs[0]
    p_h.text = tname
    p_h.alignment = PP_ALIGN.CENTER
    p_h.font.bold = True
    p_h.font.size = Pt(16)
    p_h.font.color.rgb = WHITE
    
    tb = slide_7.shapes.add_textbox(x + Inches(0.3), Inches(2.7), col_w - Inches(0.6), Inches(2.5))
    tf = tb.text_frame
    tf.word_wrap = True
    
    for idx, fld in enumerate(fields):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = f"•  {fld}"
        p.font.size = Pt(13)
        p.font.color.rgb = NAVY if "(PK)" in fld else SLATE
        p.font.bold = True if "(PK)" in fld else False
        p.space_after = Pt(8)

# Bottom Explanation Banner
banner = add_card(slide_7, Inches(0.8), Inches(5.7), Inches(11.7), Inches(0.9), NAVY, None)
tb_b = slide_7.shapes.add_textbox(Inches(1.0), Inches(5.8), Inches(11.3), Inches(0.7))
tf_b = tb_b.text_frame
tf_b.word_wrap = True
p_b = tf_b.paragraphs[0]
p_b.text = "Relational Efficiency: SQL database schema ensures seamless data retrieval, updating, and relational integrity between job applications, users, and available job postings."
p_b.font.size = Pt(13)
p_b.font.color.rgb = RGBColor(226, 232, 240)

# -----------------------------------------------------------------------------
# SLIDE 8: KEY FEATURES & ADVANTAGES
# -----------------------------------------------------------------------------
slide_8 = prs.slides.add_slide(blank_layout)
set_slide_background(slide_8)
add_header(slide_8, "Key Features & Advantages")

# Left Box: Key Features
add_card(slide_8, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
tb = slide_8.shapes.add_textbox(Inches(1.1), Inches(2.1), Inches(5.0), Inches(4.2))
tf = tb.text_frame
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "Core Features"
p.font.bold = True
p.font.size = Pt(20)
p.font.color.rgb = TEAL
p.space_after = Pt(14)

feats = [
    "User Registration & Secure Login",
    "Dynamic Job Search & Advanced Filtering",
    "Detailed Job Description Views",
    "Instant Online Application Submission",
    "Employer Job Posting & Management",
    "Structured SQL Relational Storage"
]
for ft in feats:
    p = tf.add_paragraph()
    p.text = f"★  {ft}"
    p.font.size = Pt(13)
    p.font.color.rgb = SLATE
    p.space_after = Pt(8)

# Right Box: Advantages
add_card(slide_8, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8))
tb = slide_8.shapes.add_textbox(Inches(7.1), Inches(2.1), Inches(5.1), Inches(4.2))
tf = tb.text_frame
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "Key Advantages"
p.font.bold = True
p.font.size = Pt(20)
p.font.color.rgb = ACCENT_GREEN
p.space_after = Pt(14)

advs = [
    "Saves Time: Fast, automated recruitment cycle.",
    "Easy to Use: Clean, simple, and responsive UI.",
    "Reduces Paperwork: Digitized document handling.",
    "Centralized: Single platform for jobs and applicants.",
    "Organized: Structured database for candidate management."
]
for ad in advs:
    p = tf.add_paragraph()
    p.text = f"✔  {ad}"
    p.font.size = Pt(13)
    p.font.color.rgb = SLATE
    p.space_after = Pt(10)

# -----------------------------------------------------------------------------
# SLIDE 9: FUTURE SCOPE & PROJECT OUTCOME
# -----------------------------------------------------------------------------
slide_9 = prs.slides.add_slide(blank_layout)
set_slide_background(slide_9)
add_header(slide_9, "Future Enhancements & Outcome")

# Left Column: Future Scope
add_card(slide_9, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
tb = slide_9.shapes.add_textbox(Inches(1.1), Inches(2.1), Inches(5.0), Inches(4.2))
tf = tb.text_frame
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "Future Scope"
p.font.bold = True
p.font.size = Pt(20)
p.font.color.rgb = ACCENT_PURPLE
p.space_after = Pt(14)

fscopes = [
    "AI-based Job Recommendations",
    "Resume Upload & Automated Parser",
    "Email / SMS Instant Notifications",
    "Integrated Online Video Interviewing",
    "Mobile Application (Android/iOS)",
    "Company Reviews & Ratings System",
    "AI Career Assistance Chatbot"
]
for fs in fscopes:
    p = tf.add_paragraph()
    p.text = f"🚀  {fs}"
    p.font.size = Pt(12)
    p.font.color.rgb = SLATE
    p.space_after = Pt(6)

# Right Column: Project Outcome
add_card(slide_9, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8), NAVY, None)
tb = slide_9.shapes.add_textbox(Inches(7.1), Inches(2.1), Inches(5.1), Inches(4.2))
tf = tb.text_frame
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "Project Outcome"
p.font.bold = True
p.font.size = Pt(20)
p.font.color.rgb = TEAL
p.space_after = Pt(16)

p_out = tf.add_paragraph()
p_out.text = "The Job Portal successfully provides a clean, functional web platform bridging the gap between job seekers and hiring recruiters."
p_out.font.size = Pt(15)
p_out.font.color.rgb = WHITE
p_out.space_after = Pt(16)

p_out2 = tf.add_paragraph()
p_out2.text = "It demonstrates the practical application of frontend technologies (HTML, CSS, JS) connected with structured database management (SQL) for real-world web development."
p_out2.font.size = Pt(13)
p_out2.font.color.rgb = RGBColor(226, 232, 240)

# -----------------------------------------------------------------------------
# SLIDE 10: CONCLUSION & THANK YOU
# -----------------------------------------------------------------------------
slide_10 = prs.slides.add_slide(blank_layout)
set_slide_background(slide_10)

# Left Half: Conclusion & Key Learnings
add_card(slide_10, Inches(0.8), Inches(0.8), Inches(5.6), Inches(5.9))
tb = slide_10.shapes.add_textbox(Inches(1.1), Inches(1.1), Inches(5.0), Inches(5.3))
tf = tb.text_frame
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "Conclusion"
p.font.bold = True
p.font.size = Pt(22)
p.font.color.rgb = NAVY
p.space_after = Pt(10)

p_c = tf.add_paragraph()
p_c.text = "The Job Portal project delivers an efficient recruitment platform while mastering full-stack web concepts."
p_c.font.size = Pt(13)
p_c.font.color.rgb = SLATE
p_c.space_after = Pt(20)

p_l = tf.add_paragraph()
p_l.text = "Key Project Learnings:"
p_l.font.bold = True
p_l.font.size = Pt(16)
p_l.font.color.rgb = TEAL
p_l.space_after = Pt(10)

learnings = [
    "Frontend Layout & Web Design",
    "JavaScript Logic & Validation",
    "SQL Database Architecture",
    "UI/UX & User Navigation Flow",
    "Problem Solving & Presentation Skills"
]
for lr in learnings:
    p = tf.add_paragraph()
    p.text = f"✔  {lr}"
    p.font.size = Pt(12)
    p.font.color.rgb = SLATE
    p.space_after = Pt(6)

# Right Half: Thank You Card
add_card(slide_10, Inches(6.8), Inches(0.8), Inches(5.7), Inches(5.9), NAVY, None)
tb_ty = slide_10.shapes.add_textbox(Inches(7.1), Inches(2.2), Inches(5.1), Inches(3.5))
tf_ty = tb_ty.text_frame
tf_ty.word_wrap = True

p_ty = tf_ty.paragraphs[0]
p_ty.text = "THANK YOU"
p_ty.alignment = PP_ALIGN.CENTER
p_ty.font.bold = True
p_ty.font.size = Pt(40)
p_ty.font.color.rgb = WHITE

p_q = tf_ty.add_paragraph()
p_q.text = "Any Questions?"
p_q.alignment = PP_ALIGN.CENTER
p_q.font.size = Pt(22)
p_q.font.color.rgb = TEAL
p_q.space_after = Pt(20)

p_subty = tf_ty.add_paragraph()
p_subty.text = "Job Portal Project Presentation\nBCA – 3rd Semester"
p_subty.alignment = PP_ALIGN.CENTER
p_subty.font.size = Pt(13)
p_subty.font.color.rgb = RGBColor(148, 163, 184)

# Save Presentation
prs.save("Job_Portal_Presentation.pptx")